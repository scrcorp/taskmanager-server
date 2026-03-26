"""Admin inventory router — Categories, Products, Store Inventory, Transactions, Audits.

All inventory admin endpoints consolidated in one router file.
"""

import io
from typing import Annotated
from uuid import UUID

from fastapi import APIRouter, Depends, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import check_store_access, require_permission
from app.database import get_db
from app.models.user import User
from app.schemas.inventory import (
    CategoryCreate, CategoryUpdate, CategoryTreeResponse,
    ProductCreate, ProductUpdate, ProductResponse, ProductDetailResponse,
    StoreInventoryBulkAdd, StoreInventoryUpdate, StoreInventoryResponse,
    TransactionCreate, TransactionResponse, BulkTransactionCreate,
    AuditResponse, AuditDetailResponse,
    AuditSettingUpdate, AuditSettingResponse,
)
from app.services.inventory_service import (
    category_service, sub_unit_service, product_service,
    store_inventory_service,
    transaction_service, audit_service, audit_setting_service,
)

router = APIRouter()


# ═══════════════════════════════════════════════════
# Categories
# ═══════════════════════════════════════════════════

@router.get("/inventory/categories", response_model=list[CategoryTreeResponse])
async def list_categories(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
) -> list[CategoryTreeResponse]:
    return await category_service.list_tree(db, current_user.organization_id)


@router.post("/inventory/categories", status_code=201)
async def create_category(
    data: CategoryCreate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:create"))],
):
    return await category_service.create(db, current_user.organization_id, data)


@router.put("/inventory/categories/{category_id}")
async def update_category(
    category_id: UUID,
    data: CategoryUpdate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:update"))],
):
    return await category_service.update(db, category_id, current_user.organization_id, data)


@router.delete("/inventory/categories/{category_id}", status_code=204)
async def delete_category(
    category_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:delete"))],
) -> None:
    await category_service.delete(db, category_id, current_user.organization_id)


# ═══════════════════════════════════════════════════
# Sub Units
# ═══════════════════════════════════════════════════

@router.get("/inventory/sub-units")
async def list_sub_units(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
) -> list:
    return await sub_unit_service.list_sub_units(db, current_user.organization_id)


@router.post("/inventory/sub-units", status_code=201)
async def create_sub_unit(
    data: dict,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:create"))],
) -> dict:
    return await sub_unit_service.create(
        db, current_user.organization_id, data.get("name", ""), data.get("code"),
    )


@router.put("/inventory/sub-units/{sub_unit_id}")
async def update_sub_unit(
    sub_unit_id: UUID,
    data: dict,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:update"))],
) -> dict:
    return await sub_unit_service.update(
        db, sub_unit_id, current_user.organization_id, data.get("name", ""),
    )


@router.delete("/inventory/sub-units/{sub_unit_id}", status_code=204)
async def delete_sub_unit(
    sub_unit_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:delete"))],
) -> None:
    await sub_unit_service.delete(db, sub_unit_id, current_user.organization_id)


# ═══════════════════════════════════════════════════
# Products
# ═══════════════════════════════════════════════════

@router.get("/inventory/products")
async def list_products(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
    keyword: str | None = None,
    search_field: str | None = None,
    category_id: str | None = None,
    is_active: bool | None = None,
    page: int = 1,
    per_page: int = 20,
) -> dict:
    items, total = await product_service.list_products(
        db, current_user.organization_id, keyword, search_field, category_id, is_active, page, per_page
    )
    return {"items": items, "total": total, "page": page, "per_page": per_page}


@router.post("/inventory/products", response_model=ProductResponse, status_code=201)
async def create_product(
    data: ProductCreate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:create"))],
) -> ProductResponse:
    return await product_service.create_product(db, current_user.organization_id, data, current_user.id)


@router.get("/inventory/products/generate-code")
async def preview_product_code(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
) -> dict:
    """Preview what auto-generated product code would be."""
    code = await product_service.preview_code(db, current_user.organization_id)
    return {"code": code}


@router.get("/inventory/products/excel-template")
async def download_excel_template(
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
) -> StreamingResponse:
    """Download Excel template for bulk product import."""
    try:
        import openpyxl
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(status_code=500, detail="openpyxl not installed on server")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"

    # Headers
    headers = [
        "name *", "code", "category", "subcategory", "sub_unit", "sub_unit_ratio **",
        "description", "store_code **", "min_quantity", "initial_quantity", "is_frequent",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = openpyxl.styles.Font(bold=True)

    # Instructions row
    instructions = [
        "Product name (required)",
        "Leave empty for auto-generate (P-XXXXXXXX)",
        "Category name (auto-created if new)",
        "Subcategory name (optional)",
        "e.g. box, pack, case (optional)",
        "** required if sub_unit is set (ea per sub_unit)",
        "Product description (optional)",
        "** required to link product to store (set store code in admin first)",
        "Min stock alert threshold (default 0)",
        "Starting stock quantity (default 0)",
        "yes/no (default no)",
    ]
    for col, inst in enumerate(instructions, 1):
        cell = ws.cell(row=2, column=col, value=inst)
        cell.font = openpyxl.styles.Font(italic=True, color="888888")

    # Example rows
    examples = [
        ["Whole Milk (1L)", "", "Beverages", "Dairy", "case", 12, "Fresh pasteurized milk", "DT", 10, 24, "yes"],
        ["Paper Cups (12oz)", "", "Supplies", "Packaging", "sleeve", 50, "Double-wall insulated", "DT", 100, 200, "no"],
        ["Croissant", "", "Food", "", "", "", "Butter croissant, baked daily", "DT", 10, 8, "yes"],
        ["Dish Soap (1L)", "P-CUSTOM01", "Supplies", "Cleaning", "", "", "", "WS", 3, 5, "no"],
    ]
    for row_idx, example in enumerate(examples, 3):
        for col_idx, val in enumerate(example, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = openpyxl.styles.Font(color="4472C4")

    # Note row
    note_row = len(examples) + 4
    ws.cell(row=note_row, column=1, value="DELETE EXAMPLE ROWS ABOVE BEFORE UPLOADING").font = openpyxl.styles.Font(bold=True, color="FF0000")
    ws.cell(row=note_row + 1, column=1, value="* = always required. ** = conditionally required (see instruction row). Leave optional fields empty or use '-' to skip.").font = openpyxl.styles.Font(italic=True, color="888888")
    ws.cell(row=note_row + 2, column=1, value="If product code already exists, it will not create a new product — you can choose to link it to a store instead.").font = openpyxl.styles.Font(italic=True, color="888888")

    # Auto-width
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=inventory_import_template.xlsx"},
    )


@router.post("/inventory/products/import")
async def import_products_from_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_permission("inventory:create")),
) -> dict:
    """Import products from uploaded Excel file."""
    try:
        import openpyxl
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(status_code=500, detail="openpyxl not installed on server")

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return {"error": "Empty file"}

    # Normalize headers
    headers = [str(h or "").strip().lower().replace(" *", "").replace("*", "") for h in header_row]

    parsed_rows = []
    for row in rows_iter:
        row_dict = {}
        for i, val in enumerate(row):
            if i < len(headers):
                key = headers[i]
                str_val = str(val).strip() if val is not None else ""
                if str_val in ("-", ""):
                    str_val = ""
                row_dict[key] = str_val
        # Skip instruction/empty rows
        if not row_dict.get("name") or row_dict["name"].startswith("Product name") or row_dict["name"].startswith("DELETE"):
            continue
        parsed_rows.append(row_dict)

    result = await product_service.import_from_excel(
        db, current_user.organization_id, parsed_rows, current_user.id
    )
    return result


# ── Product CRUD with path params (MUST be after /generate-code, /excel-template, /import) ──

@router.get("/inventory/products/{product_id}", response_model=ProductDetailResponse)
async def get_product(
    product_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
) -> ProductDetailResponse:
    return await product_service.get_product(db, product_id, current_user.organization_id)


@router.put("/inventory/products/{product_id}", response_model=ProductResponse)
async def update_product(
    product_id: UUID,
    data: ProductUpdate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:update"))],
) -> ProductResponse:
    return await product_service.update_product(db, product_id, current_user.organization_id, data)


@router.delete("/inventory/products/{product_id}", status_code=204)
async def deactivate_product(
    product_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:delete"))],
) -> None:
    await product_service.deactivate_product(db, product_id, current_user.organization_id)


# ═══════════════════════════════════════════════════
# Store Inventory
# ═══════════════════════════════════════════════════

@router.get("/stores/{store_id}/inventory")
async def list_store_inventory(
    store_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
    keyword: str | None = None,
    search_field: str | None = None,
    status: str | None = None,
    is_frequent: bool | None = None,
    page: int = 1,
    per_page: int = 50,
) -> dict:
    await check_store_access(db, current_user, store_id)
    items, total = await store_inventory_service.list_inventory(
        db, store_id, current_user.organization_id,
        keyword, search_field, status, is_frequent, page, per_page,
    )
    return {"items": items, "total": total, "page": page, "per_page": per_page}


@router.get("/stores/{store_id}/inventory/summary")
async def get_store_inventory_summary(
    store_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
) -> dict:
    await check_store_access(db, current_user, store_id)
    return await store_inventory_service.get_summary(db, store_id, current_user.organization_id)


@router.post("/stores/{store_id}/inventory", status_code=201)
async def bulk_add_store_inventory(
    store_id: UUID,
    data: StoreInventoryBulkAdd,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:create"))],
) -> list[StoreInventoryResponse]:
    await check_store_access(db, current_user, store_id)
    return await store_inventory_service.bulk_add(
        db, store_id, current_user.organization_id, data, current_user.id
    )


@router.put("/stores/{store_id}/inventory/{item_id}", response_model=StoreInventoryResponse)
async def update_store_inventory_item(
    store_id: UUID, item_id: UUID,
    data: StoreInventoryUpdate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:update"))],
) -> StoreInventoryResponse:
    await check_store_access(db, current_user, store_id)
    return await store_inventory_service.update_item(
        db, item_id, store_id, current_user.organization_id, data
    )


@router.delete("/stores/{store_id}/inventory/{item_id}", status_code=204)
async def deactivate_store_inventory_item(
    store_id: UUID, item_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:delete"))],
) -> None:
    await check_store_access(db, current_user, store_id)
    await store_inventory_service.deactivate_item(
        db, item_id, store_id, current_user.organization_id
    )


# ═══════════════════════════════════════════════════
# Transactions
# ═══════════════════════════════════════════════════

@router.get("/stores/{store_id}/inventory/transactions")
async def list_store_transactions(
    store_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
    product_id: str | None = None,
    type: str | None = None,
    page: int = 1,
    per_page: int = 20,
) -> dict:
    await check_store_access(db, current_user, store_id)
    items, total = await transaction_service.list_transactions(
        db, store_id, current_user.organization_id, product_id, type, page, per_page
    )
    return {"items": items, "total": total, "page": page, "per_page": per_page}


@router.post("/stores/{store_id}/inventory/{item_id}/transactions", response_model=TransactionResponse, status_code=201)
async def create_transaction(
    store_id: UUID, item_id: UUID,
    data: TransactionCreate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:create"))],
) -> TransactionResponse:
    await check_store_access(db, current_user, store_id)
    return await transaction_service.create_transaction(
        db, item_id, store_id, current_user.organization_id, data, current_user.id
    )


@router.post("/stores/{store_id}/inventory/bulk-stock-in", status_code=201)
async def bulk_stock_in(
    store_id: UUID,
    data: BulkTransactionCreate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:create"))],
) -> list[TransactionResponse]:
    await check_store_access(db, current_user, store_id)
    data.type = "stock_in"
    return await transaction_service.bulk_transaction(
        db, store_id, current_user.organization_id, data, current_user.id
    )


@router.post("/stores/{store_id}/inventory/bulk-stock-out", status_code=201)
async def bulk_stock_out(
    store_id: UUID,
    data: BulkTransactionCreate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:create"))],
) -> list[TransactionResponse]:
    await check_store_access(db, current_user, store_id)
    data.type = "stock_out"
    return await transaction_service.bulk_transaction(
        db, store_id, current_user.organization_id, data, current_user.id
    )


# ═══════════════════════════════════════════════════
# Audits
# ═══════════════════════════════════════════════════

@router.get("/stores/{store_id}/inventory/audits")
async def list_audits(
    store_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
    page: int = 1,
    per_page: int = 20,
) -> dict:
    await check_store_access(db, current_user, store_id)
    items, total = await audit_service.list_audits(
        db, store_id, current_user.organization_id, page, per_page
    )
    return {"items": items, "total": total, "page": page, "per_page": per_page}


@router.get("/stores/{store_id}/inventory/audits/{audit_id}", response_model=AuditDetailResponse)
async def get_audit(
    store_id: UUID, audit_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
) -> AuditDetailResponse:
    await check_store_access(db, current_user, store_id)
    return await audit_service.get_audit_detail(db, audit_id, store_id, current_user.organization_id)


# ═══════════════════════════════════════════════════
# Audit Settings
# ═══════════════════════════════════════════════════

@router.get("/stores/{store_id}/inventory/audit-settings", response_model=AuditSettingResponse)
async def get_audit_settings(
    store_id: UUID,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:read"))],
) -> AuditSettingResponse:
    await check_store_access(db, current_user, store_id)
    return await audit_setting_service.get_setting(db, store_id, current_user.organization_id)


@router.put("/stores/{store_id}/inventory/audit-settings", response_model=AuditSettingResponse)
async def update_audit_settings(
    store_id: UUID,
    data: AuditSettingUpdate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_permission("inventory:update"))],
) -> AuditSettingResponse:
    await check_store_access(db, current_user, store_id)
    return await audit_setting_service.update_setting(db, store_id, current_user.organization_id, data)
